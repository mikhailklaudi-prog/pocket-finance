import asyncio
import datetime
import io
import json
import logging
import socket
import sqlite3
import urllib.request
from aiohttp import web
from aiogram import Bot, Dispatcher, types, F
from aiogram.client.session.aiohttp import AiohttpSession
from aiogram.filters import Command, CommandStart
from aiogram.types import InlineKeyboardMarkup, InlineKeyboardButton, WebAppInfo, BufferedInputFile
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

logging.basicConfig(level=logging.INFO, format="%(asctime)s - %(levelname)s - %(message)s")

# IPv4 фикс для устранения таймаутов в Windows
_orig_getaddrinfo = socket.getaddrinfo
def _ipv4_only_getaddrinfo(*args, **kwargs):
    responses = _orig_getaddrinfo(*args, **kwargs)
    return [r for r in responses if r[0] == socket.AF_INET]
socket.getaddrinfo = _ipv4_only_getaddrinfo

BOT_TOKEN = "7582044537:AAH09vAKGEz5LUVs-cc-KnRXBeYx3TrFKIg"
BASE_URL = "https://d9xndp1j-3000.euw.devtunnels.ms/index.html"
DB_NAME = "finance.db"

# Подключение прокси VPN
sys_proxies = urllib.request.getproxies()
proxy_url = sys_proxies.get("https") or sys_proxies.get("http")
if proxy_url:
    print(f"[+] VPN-прокси активен: {proxy_url}")
    session = AiohttpSession(proxy=proxy_url)
    bot = Bot(token=BOT_TOKEN, session=session)
else:
    bot = Bot(token=BOT_TOKEN)

dp = Dispatcher()

# --- База данных SQLite ---
def init_db():
    conn = sqlite3.connect(DB_NAME)
    cur = conn.cursor()
    cur.execute("""
        CREATE TABLE IF NOT EXISTS transactions (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            user_id INTEGER NOT NULL,
            type TEXT NOT NULL,
            category TEXT NOT NULL,
            amount REAL NOT NULL,
            note TEXT,
            date_str TEXT,
            created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP
        )
    """)
    cur.execute("""
        CREATE TABLE IF NOT EXISTS users (
            user_id INTEGER PRIMARY KEY,
            first_name TEXT,
            username TEXT
        )
    """)
    try:
        cur.execute("ALTER TABLE transactions ADD COLUMN date_str TEXT")
    except sqlite3.OperationalError:
        pass
    conn.commit()
    conn.close()

def db_register_user(user_id: int, first_name: str, username: str):
    conn = sqlite3.connect(DB_NAME)
    cur = conn.cursor()
    cur.execute(
        "INSERT OR REPLACE INTO users (user_id, first_name, username) VALUES (?, ?, ?)",
        (user_id, first_name, username)
    )
    conn.commit()
    conn.close()

def db_add_transaction(user_id: int, t_type: str, category: str, amount: float, note: str, date_str: str):
    conn = sqlite3.connect(DB_NAME)
    cur = conn.cursor()
    cur.execute(
        "INSERT INTO transactions (user_id, type, category, amount, note, date_str) VALUES (?, ?, ?, ?, ?, ?)",
        (user_id, t_type, category, amount, note, date_str)
    )
    conn.commit()
    conn.close()

def db_delete_transaction(user_id: int, transaction_id: int):
    conn = sqlite3.connect(DB_NAME)
    cur = conn.cursor()
    cur.execute(
        "DELETE FROM transactions WHERE id = ? AND (user_id = ? OR user_id = 0)",
        (transaction_id, user_id)
    )
    conn.commit()
    conn.close()

def db_get_data(user_id: int):
    conn = sqlite3.connect(DB_NAME)
    cur = conn.cursor()

    target_id = user_id
    if target_id == 0:
        cur.execute("SELECT user_id FROM transactions WHERE user_id != 0 ORDER BY id DESC LIMIT 1")
        row = cur.fetchone()
        if row:
            target_id = row[0]

    cur.execute("""
        SELECT 
            COALESCE(SUM(CASE WHEN type = 'income' THEN amount ELSE 0 END), 0),
            COALESCE(SUM(CASE WHEN type = 'expense' THEN amount ELSE 0 END), 0),
            COUNT(id)
        FROM transactions WHERE user_id = ? OR user_id = 0
    """, (target_id,))
    inc, exp, count = cur.fetchone()

    cur.execute("""
        SELECT id, type, category, amount, note, date_str 
        FROM transactions 
        WHERE user_id = ? OR user_id = 0
        ORDER BY id DESC LIMIT 50
    """, (target_id,))
    rows = cur.fetchall()
    conn.close()

    transactions = [
        {"id": r[0], "type": r[1], "category": r[2], "amount": r[3], "note": r[4], "date": r[5] or ""}
        for r in rows
    ]

    return {
        "balance": inc - exp,
        "income": inc,
        "expense": exp,
        "count": count,
        "transactions": transactions
    }

def db_reset(user_id: int):
    conn = sqlite3.connect(DB_NAME)
    cur = conn.cursor()
    cur.execute("DELETE FROM transactions WHERE user_id = ? OR user_id = 0", (user_id,))
    conn.commit()
    conn.close()

# --- Формирование Excel-таблицы ---
def create_excel_report(user_id: int) -> io.BytesIO:
    conn = sqlite3.connect(DB_NAME)
    cur = conn.cursor()

    target_id = user_id
    if target_id == 0:
        cur.execute("SELECT user_id FROM transactions WHERE user_id != 0 ORDER BY id DESC LIMIT 1")
        row = cur.fetchone()
        if row:
            target_id = row[0]

    cur.execute("""
        SELECT id, created_at, type, category, amount, note 
        FROM transactions 
        WHERE user_id = ? OR user_id = 0
        ORDER BY id ASC
    """, (target_id,))
    rows = cur.fetchall()
    conn.close()

    wb = Workbook()
    ws = wb.active
    ws.title = "Операции"

    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    regular_font = Font(name="Segoe UI", size=10)
    inc_font = Font(name="Segoe UI", size=10, bold=True, color="16A34A")
    exp_font = Font(name="Segoe UI", size=10, bold=True, color="DC2626")
    border_side = Side(style='thin', color="CBD5E1")
    cell_border = Border(left=border_side, right=border_side, top=border_side, bottom=border_side)

    headers = ["№", "Дата и время", "Тип", "Категория", "Сумма (₽)", "Заметка"]
    ws.append(headers)

    for col_idx in range(1, len(headers) + 1):
        c = ws.cell(row=1, column=col_idx)
        c.fill = header_fill
        c.font = header_font
        c.alignment = Alignment(horizontal="center", vertical="center")

    for i, r in enumerate(rows, start=1):
        t_type = "Доход" if r[2] == "income" else ("Заметка" if r[2] == "note" else "Расход")
        sign = "+" if r[2] == "income" else ("" if r[2] == "note" else "-")
        amount_val = f"{sign}{r[4]:,.2f} ₽" if r[2] != "note" else "—"

        row_data = [i, r[1], t_type, r[3], amount_val, r[5] or ""]
        ws.append(row_data)

        current_row = ws.max_row
        for col_idx in range(1, len(row_data) + 1):
            c = ws.cell(row=current_row, column=col_idx)
            c.border = cell_border
            if col_idx == 5:
                c.font = inc_font if r[2] == "income" else (exp_font if r[2] == "expense" else regular_font)
                c.alignment = Alignment(horizontal="right")
            elif col_idx in (1, 2, 3):
                c.font = regular_font
                c.alignment = Alignment(horizontal="center")
            else:
                c.font = regular_font

    for col in ws.columns:
        col_letter = get_column_letter(col[0].column)
        max_len = max(len(str(cell.value or '')) for cell in col)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf

# --- Веб-сервер API ---
async def serve_index(request):
    return web.FileResponse("index.html")

async def serve_app_js(request):
    return web.FileResponse("app.js")

async def api_get_data(request):
    user_id = int(request.query.get("user_id", 0))
    return web.json_response(db_get_data(user_id), headers={"Access-Control-Allow-Origin": "*"})

async def api_add_transaction(request):
    body = await request.json()
    user_id = int(body.get("user_id", 0))
    db_add_transaction(
        user_id,
        body.get("type", "expense"),
        body.get("category", "Разное"),
        float(body.get("amount", 0)),
        body.get("note", ""),
        body.get("date", "")
    )
    return web.json_response(db_get_data(user_id), headers={"Access-Control-Allow-Origin": "*"})

async def api_delete_transaction(request):
    body = await request.json()
    user_id = int(body.get("user_id", 0))
    t_id = int(body.get("id", 0))
    db_delete_transaction(user_id, t_id)
    return web.json_response(db_get_data(user_id), headers={"Access-Control-Allow-Origin": "*"})

async def api_reset_data(request):
    body = await request.json()
    user_id = int(body.get("user_id", 0))
    db_reset(user_id)
    return web.json_response({"ok": True}, headers={"Access-Control-Allow-Origin": "*"})

# --- Бот и обработчики команд ---
def main_kb(user_id: int):
    return InlineKeyboardMarkup(
        inline_keyboard=[
            [InlineKeyboardButton(text="Открыть кошелёк 💳", web_app=WebAppInfo(url=f"{BASE_URL}?user_id={user_id}"))],
            [
                InlineKeyboardButton(text="📊 Сводка", callback_data="refresh_report"),
                InlineKeyboardButton(text="📥 Скачать Excel", callback_data="export_excel")
            ]
        ]
    )

@dp.message(CommandStart())
async def cmd_start(message: types.Message):
    user_id = message.from_user.id
    db_register_user(user_id, message.from_user.first_name, message.from_user.username)
    text = (
        f"Привет, {message.from_user.first_name or 'друг'}! 👋\n\n"
        "Добро пожаловать в <b>Pocket Finance</b>.\n\n"
        "💳 Открывай кошелёк кнопкой ниже для учёта финансов.\n"
        "📊 /report — баланс и последние записи.\n"
        "📥 /export — выгрузка отчёта в Excel-таблицу."
    )
    await message.answer(text, reply_markup=main_kb(user_id), parse_mode="HTML")

@dp.message(Command("report"))
async def cmd_report(message: types.Message):
    await send_report(message.chat.id, message.answer)

@dp.message(Command("export"))
async def cmd_export(message: types.Message):
    await send_excel(message.chat.id)

@dp.callback_query(F.data == "refresh_report")
async def cb_refresh_report(call: types.CallbackQuery):
    await send_report(call.message.chat.id, call.message.edit_text)
    await call.answer()

@dp.callback_query(F.data == "export_excel")
async def cb_export(call: types.CallbackQuery):
    await send_excel(call.message.chat.id)
    await call.answer()

async def send_report(chat_id: int, send_func):
    data = db_get_data(chat_id)
    text = (
        "📊 <b>Финансовый отчёт (SQLite)</b>\n"
        "━━━━━━━━━━━━━━━━━━━━\n"
        f"💰 <b>Общий баланс:</b> {data['balance']:,.0f} ₽\n"
        f"📈 <b>Доходы:</b> +{data['income']:,.0f} ₽\n"
        f"📉 <b>Расходы:</b> -{data['expense']:,.0f} ₽\n"
        f"📝 <b>Всего записей:</b> {data['count']}\n"
    )
    if data["transactions"]:
        text += "━━━━━━━━━━━━━━━━━━━━\n<b>Последние операции:</b>\n"
        for t in data["transactions"][:5]:
            icon = "🟢" if t["type"] == "income" else ("🟣" if t["type"] == "note" else "🔴")
            sign = "+" if t["type"] == "income" else ("" if t["type"] == "note" else "-")
            amt = f"<b>{sign}{t['amount']:,.0f} ₽</b>" if t["type"] != "note" else "<i>заметка</i>"
            text += f"{icon} {t['category']} — {amt} <i>({t['note']})</i>\n"
    else:
        text += "━━━━━━━━━━━━━━━━━━━━\n<i>Записей в базе пока нет.</i>\n"

    try:
        await send_func(text, reply_markup=main_kb(chat_id), parse_mode="HTML")
    except Exception:
        pass

# Прямая отправка документа через bot.send_document
async def send_excel(chat_id: int):
    buf = create_excel_report(chat_id)
    today = datetime.date.today().strftime("%d.%m.%Y")
    file = BufferedInputFile(buf.getvalue(), filename=f"PocketFinance_{today}.xlsx")
    await bot.send_document(
        chat_id=chat_id,
        document=file,
        caption="📥 <b>Твой финансовый отчёт готов!</b>",
        parse_mode="HTML"
    )

async def main():
    init_db()

    app = web.Application()
    app.router.add_get("/", serve_index)
    app.router.add_get("/index.html", serve_index)
    app.router.add_get("/app.js", serve_app_js)
    app.router.add_get("/api/data", api_get_data)
    app.router.add_post("/api/add", api_add_transaction)
    app.router.add_post("/api/delete", api_delete_transaction)
    app.router.add_post("/api/reset", api_reset_data)

    runner = web.AppRunner(app)
    await runner.setup()
    site = web.TCPSite(runner, "0.0.0.0", 3000)
    await site.start()
    print(">>> Встроенный веб-сервер слушает порт 3000")

    await bot.delete_webhook(drop_pending_updates=True)
    print(">>> Бот запущен и слушает Telegram...")
    await dp.start_polling(bot)

if __name__ == "__main__":
    try:
        asyncio.run(main())
    except KeyboardInterrupt:
        print("\n>>> Бот выключен.")